"""
Release:  V2 - XX/XX/2025
Internal: V2
Last Updated: 06/30/2025
"""
import subprocess
import sys

def install(package):
    """Automatically runs the Python pip install command to download necessary external packages"""
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

# If user does not have any packages installed, install them automatically
try:
    from tkinter import Tk
    import os
    import tkinter.filedialog
    import olefile
    from docx import Document
    import logging
    import re
    import fitz
except:
    install("python-docx")
    install("olefile")
    install("PyMuPDF")
    from docx import Document
    import olefile
    import fitz

def get_folder(message):
    """Opens a window to select a file directory"""
    Tk().withdraw()
    return tkinter.filedialog.askdirectory(title=message)

def convert_scalc_to_txt(folder):
    """Converts all Settings Basis documents into text files for data parsing"""
    basis_count = 0
    for scalc in os.listdir(folder):
        lower_name = scalc.lower()
        currentpath = os.path.join(folder, scalc)
        if ("basis" in lower_name) and ("settings" in lower_name) and ("network" not in lower_name):
            document_name = os.path.splitext(scalc)[0]
            txt_path = os.path.join(folder, document_name + ".txt")
            if lower_name.endswith(".docx"):
                basis_count += 1
                logging.info(f"Converting {scalc}...")
                doc = Document(currentpath)
                with open(txt_path, 'w', encoding='utf-8') as file:
                    for table in doc.tables:
                        for row in table.rows:
                            row_text = '\t'.join(cell.text.strip() for cell in row.cells)
                            file.write(row_text + '\n')
            elif lower_name.endswith(".pdf"):
                # docx_path = os.path.join(folder, document_name + ".docx")
                # if not os.path.exists(docx_path):
                #     pdf_now_doc = Converter(currentpath)
                #     pdf_now_doc.convert(docx_path, start=0, end=None)
                #     pdf_now_doc.close()
                #     try:
                #         doc = Document(docx_path)
                #         all_text = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
                #         with open(txt_path, 'w', encoding='utf-8') as f:
                #             f.write("\n".join(all_text))
                #         basis_count += 1
                #     except Exception as e:
                #         logging.error(f"Failed to read generated DOCX: {docx_path} — {e}")
                #         basis_count -= 1
                #     finally:
                #         if os.path.exists(docx_path):
                #             os.remove(docx_path)
                logging.info(f"Converting {scalc}...")
                currentpath = os.path.join(folder, scalc)
                document_name = os.path.splitext(scalc)[0]
                txt_path = os.path.join(folder, document_name + ".txt")
                setting_lines = []
                with fitz.open(currentpath) as pdf:
                    text = ""
                    for page in pdf:
                        text += page.get_text()
                        for line in text.splitlines():
                            if ":=" in line:
                                setting_lines.append(line)
                if setting_lines:
                    basis_count += 1
                    with open(txt_path, 'w', encoding='utf-8') as file:
                        file.write("\n".join(setting_lines))
    return basis_count

def clean_txts(folder):
    """Cleans the Settings Basis text files by parsing out the settings values in a nested dictionary."""
    # key: device name
    # value: dictionary of settings
    txt_files = {}
    for txt_file in os.listdir(folder):
        if ("basis" in txt_file.lower()) and ("settings" in txt_file.lower()) and (".txt" in txt_file.lower()):
            logging.info(f"Cleaning {txt_file}...")
            currentpath = os.path.join(folder, txt_file)
            if "_Relay" in txt_file:
                device_name = txt_file.split("_Relay")[0]
            else:
                device_name = txt_file.split("_Settings")[0]
            txt_files[device_name] = {}
            with open(currentpath, "r", encoding="utf-8") as file:
                for line in file:
                    if ":=" in line:
                        key = line.split(":=")[0].split("\t")[-1].split()[-1].strip()
                        value = line.split(":=")[1].strip().split("\t")[0].split("\r")[0]
                        if 'not used' in value.lower():
                            value = value.split(" ")[0]
                        txt_files[device_name][key] = value
    return txt_files

def delete_txts(folder):
    """Deletes the newly created text files from Settings Basis docx. They are no longer needed after program concludes."""
    paths = []
    for txt_file in os.listdir(folder):
        if ("basis" in txt_file.lower()) and ("settings" in txt_file.lower()) and (".txt" in txt_file.lower()):
            paths.append(os.path.join(folder, txt_file))
    for file in paths:
        os.remove(file)

def determine_group(txt_file):
    if re.search(r"set_s\d.txt", txt_file.lower()) or re.search(r"set_\d.txt", txt_file.lower()):
        digit = re.findall(r'\d', txt_file)
        return digit[-1]
    else:
        return -1

def clean_settings_files(folder):
    """Cleans the rdb settings text files by parsing out the settings values in a nested dictionary. """
    # key: device name
    # value: dictionary of settings
    settings_files = {}
    for settings_file in os.listdir(folder):
        currentpath = folder + "\\" + settings_file
        if (".rdb" in settings_file.lower()) and ("settings" in settings_file.lower()):
            if "_Relay" in settings_file:
                device_name = settings_file.split("_Relay")[0]
            else:
                device_name = settings_file.split("_Settings")[0]
            logging.info(f"Cleaning {settings_file}...")
            settings_files[device_name] = {}
            if olefile.isOleFile(currentpath):
                ole = olefile.OleFileIO(currentpath)
                print(ole.listdir())
                for entry in ole.listdir():
                    txt_file = "/".join(entry)
                    if ".txt" in txt_file.lower():
                        if int(determine_group(txt_file.split("/")[-1])) > 1:
                            continue
                        with ole.openstream(entry) as stream:
                            content = stream.read().decode('utf-8', errors='ignore')
                            content = content.split("\r")
                            for line in content:
                                if ',"' in line or ",'" in line:
                                    if ',"' in line:
                                        l = line.split(',"')
                                    else:
                                        l = line.split(",'")
                                    key = l[0].strip()
                                    if ',""' in line or ",''" in line:
                                        value = ''
                                    else:
                                        value = l[1].strip().strip('"').strip("'")
                                    settings_files[device_name][key] = value
    return settings_files



def compare(txt_files, settings_files):
    """Compares the dictionaries acquired from previous data cleaning. Settings Basis points are compared to Settings file points."""
    output_summary = "Settings Basis vs. Settings File Comparison Summary\n\n"
    # Iterate over all settings basis text files
    for device, points in txt_files.items():
        print(f"{device:<40}{'Attribute':<20}{'Settings Basis':<40}{'Settings File':<40}")
        print(f"{'-':<40}{'-':<20}{'-':<40}{'-':<40}")
        mismatch = False
        # Find matching settings file
        matching_rdb = settings_files.get(device)
        if matching_rdb:
            # Add device name
            output_summary += f"***** {device} *****\n\n"
            # Iterate over all points in an individual settings basis
            for key, value in points.items():
                if "–" in key or "-" in key:
                    output_summary += f"{key} is listed in Settings Basis as a range of values. Please check manually.\n\n"
                    mismatch = True
                    continue
                # Compare setting value in settings basis to the settings file value
                matching_point = settings_files[device].get(key)
                try:
                    print(f"{'':<40}{key:<20}{value:<40}{matching_point:<40}")
                except:
                    print(f"{'':<40}{key:<20}{value:<40}{'':<40}")
                if not matching_point:
                    # Discrepancy was found
                    output_summary += f"{key} was not found in settings file.\n\n"
                    mismatch = True
                elif value != matching_point:
                    # Ensure numerical values are the same
                    try:
                        if float(value) != float(matching_point):
                            # Discrepancy was found
                            output_summary += f"The values of {key} do not match.\n"
                            output_summary += f"Settings Basis: {key} := {value}\n"
                            output_summary += f"Settings File: {key} := {matching_point}\n\n"
                            mismatch = True
                    except ValueError:
                        # Discrepancy was found
                        output_summary += f"The values of {key} do not match.\n"
                        output_summary += f"Settings Basis: {key} := {value}\n"
                        output_summary += f"Settings File: {key} := {matching_point}\n\n"
                        mismatch = True
        else:
            output_summary += f"Settings file for {device} was not found.\n\n"
        if not mismatch and matching_rdb:
            # Discrepancy was not found
            output_summary += "No discrepancies found!\n\n"
        print("\n")
    return output_summary

def main():
    # Select directory for Settings files + Settings Basis
    folder = get_folder("Select the directory containing your Settings Basis document(s) & Settings File(s).")
    if not folder:
        tkinter.messagebox.showerror("Error", "No valid directory selected.")
        return 1
    
    logging.info("Starting .docx .txt conversion...")
    
    # Process Settings Basis documents
    if convert_scalc_to_txt(folder) < 1:
        logging.info("No Settings Basis documents found.")
        tkinter.messagebox.showinfo("", f"No Settings Basis documents to compare.")
        return 2
    else:
        txt_files = clean_txts(folder)

        # Process rdb files
        settings_files = clean_settings_files(folder)
        if len(settings_files) < 1:
            logging.info("No rdb files found.")
            tkinter.messagebox.showinfo("", f"No rdb files to compare.")
            return 2
        
        notes = compare(txt_files, settings_files)

        # Remove new text files
        delete_txts(folder)

        # Output notes
        output_file = "Settings Comparison Summary.txt"
        output_file = folder + "/" + output_file
        with open(output_file, 'w') as ofile:
            ofile.write(notes)
        # Display popup window with completion message
        logging.info("Comparison Complete.")
        tkinter.messagebox.showinfo("", f"Comparison Notes saved at {output_file}")
        return 0

if __name__ == "__main__":
    main()
